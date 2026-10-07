"""Build the editable paper from the saved, audited research outputs.

Requires python-docx and pandas. Render with the supplied document workflow.
"""
from pathlib import Path
import json
from xml.sax.saxutils import escape
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "paper"
OUT.mkdir(exist_ok=True)
s = pd.read_csv(ROOT / "results/summary.csv", index_col=0)
grid = pd.read_csv(ROOT / "results/sensitivity.csv")
annual = pd.read_csv(ROOT / "results/annual_returns.csv", index_col=0)
periods = pd.read_csv(ROOT / "results/subperiods.csv")
manifest = json.loads((ROOT / "data/manifest.json").read_text())
run = json.loads((ROOT / "results/run_manifest.json").read_text())
audit = json.loads((ROOT / "results/audit.json").read_text())
assert audit["status"] == "PASS"
r = s.loc["Rank blend"]; spy = s.loc["SPY buy and hold"]; eq = s.loc["Equal weight universe"]
pct = lambda x: f"{x*100:.2f}%"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.75)
sec.left_margin = sec.right_margin = Inches(.8)
sec.header_distance = sec.footer_distance = Inches(.3)
for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption"]:
    st = doc.styles[name]
    st.font.name = "Times New Roman"
    st.font.color.rgb = RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    rf = st.element.get_or_add_rPr().rFonts
    for attr in ["asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme", "csTheme"]:
        rf.attrib.pop(qn("w:"+attr), None)
    rf.set(qn("w:cs"), "Times New Roman")
for st in doc.styles:
    for border in list(st.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
normal = doc.styles["Normal"]
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.1
normal.paragraph_format.space_after = Pt(7)
for name,size in [("Title",23),("Subtitle",12),("Heading 1",16),("Heading 2",12)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].paragraph_format.space_before = Pt(10)
    doc.styles[name].paragraph_format.space_after = Pt(7)
doc.styles["Caption"].font.size = Pt(9)
doc.styles["Caption"].font.italic = False
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
doc.core_properties.title = "Evaluating Momentum and Low Volatility in an ETF Portfolio"
doc.core_properties.author = "Carson Chen"
doc.core_properties.subject = "Revised capstone research using historical ETF data"
md = []

def p(text, style=None):
    para = doc.add_paragraph(text, style)
    md.append(text+"\n")
    return para

def h(text, level=1):
    doc.add_heading(text,level)
    md.append("#"*level+" "+text+"\n")

def page():
    doc.add_page_break()

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = widths or [6.9/len(headers)]*len(headers)
    for col,width in zip(t.columns,widths): col.width = Inches(width)
    for j,v in enumerate(headers): t.rows[0].cells[j].text = str(v)
    for row in rows:
        cells = t.add_row().cells
        for j,v in enumerate(row): cells[j].text = str(v)
    for i,row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        dont = OxmlElement("w:cantSplit"); trpr.append(dont)
        if i==0:
            repeat=OxmlElement("w:tblHeader"); trpr.append(repeat)
        for j,cell in enumerate(row.cells):
            cell.width=Inches(widths[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement("w:tcBorders")
            for edge in ["top","left","bottom","right"]:
                item=OxmlElement("w:"+edge); item.set(qn("w:val"),"single"); item.set(qn("w:sz"),"4"); item.set(qn("w:color"),"D9D9D9"); borders.append(item)
            pr.append(borders)
            margins=OxmlElement("w:tcMar")
            for edge in ["top","left","bottom","right"]:
                item=OxmlElement("w:"+edge);item.set(qn("w:w"),"75");item.set(qn("w:type"),"dxa");margins.append(item)
            pr.append(margins)
            shade=OxmlElement("w:shd"); shade.set(qn("w:fill"),"E8EDF0" if i==0 else ("F6F6F6" if i%2==0 else "FFFFFF")); pr.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(0)
                para.paragraph_format.line_spacing=1.0
                para.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                for rr in para.runs:
                    rr.font.size=Pt(9.5);rr.font.name="Times New Roman";rr.bold=(i==0)
    md.append("| "+" | ".join(map(str,headers))+" |\n| "+" | ".join(["---"]*len(headers))+" |")
    md.extend("| "+" | ".join(map(str,row))+" |" for row in rows)
    md.append("\n")
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def picture(filename,caption):
    para=doc.add_paragraph();para.paragraph_format.space_after=Pt(1)
    para.add_run().add_picture(str(ROOT/"results"/filename),width=Inches(6.85))
    p(caption,"Caption")
    md.append(f"![{caption}](../results/{filename})\n")

p("Evaluating Momentum and Low Volatility in an ETF Portfolio", "Title")
p("Carson Chen\nRevised capstone study  |  October 2026", "Subtitle")
h("Abstract")
p(f"This study evaluates a transparent long-only ETF selection rule combining trailing momentum and low volatility. It replaces an earlier simulated-price prototype with a documented historical-data experiment. Eight ETFs are observed from 2010 through 2025, with a common evaluation period of {run['first_evaluation_date']} to {run['last_evaluation_date']}. The default portfolio selects two ETFs monthly using equally weighted percentile ranks and executes at the next trading close. Trading costs are charged at 10 basis points per dollar bought or sold. The portfolio achieves a compound annual growth rate of {pct(r.cagr)}, annualized volatility of {pct(r.annual_volatility)}, and a maximum drawdown of {pct(r.max_drawdown)}. SPY produces a higher CAGR of {pct(spy.cagr)} with higher volatility, while an equal-weight portfolio of the same eight ETFs earns {pct(eq.cagr)} with lower volatility than the factor portfolio. The results do not establish superiority over passive alternatives. The main contribution is a reproducible test of a risk-return hypothesis, including weaker outcomes, implementation controls, and parameter sensitivity.")
h("Research question")
p("Does combining recent price strength with a preference for lower volatility improve the observed return-risk trade-off relative to momentum alone, low volatility alone, and simple passive alternatives? The question is evaluated through returns, volatility, drawdowns and trading activity rather than through a single favorable chart.")
p("The motivating intuition is straightforward: momentum seeks continued strength, while the volatility component discourages unstable price histories. Neither input guarantees a future return or a floor on losses. The combination is a hypothesis about selection, not an insurance contract or a formally optimized portfolio.")
h("Scope of the revision")
p("The archived prototype used randomly generated paths labeled with eight ETF tickers. It was useful for demonstrating a workflow, but its output cannot establish historical market performance. This revision retains those eight tickers and the original idea, introduces actual adjusted prices and corrected accounting, and reports the resulting evidence. It does not claim a universe of more than 100 ETFs or attribute the present revision to the earlier project period.")

page();h("Data and research context")
h("Historical data",2)
p(f"Daily adjusted closes were downloaded from Yahoo Finance through yfinance with auto_adjust=True [1]. The snapshot contains {manifest['rows']:,} observations per ETF, from {manifest['first_observation']} through {manifest['last_observation']}. The request end date is exclusive. The download was completed on {manifest['downloaded_utc'][:10]} UTC. Data from 2010-2011 provide pre-evaluation history; reported portfolio performance begins in January 2012. The sample ends with the last full calendar year, 2025.")
table(["Ticker","Exposure represented"],[["SPY","US large-cap equities"],["QQQ","Nasdaq-100 equities"],["IWM","US small-cap equities"],["EFA","Developed equities outside the US and Canada"],["VNQ","US listed real estate"],["GLD","Gold"],["TLT","Long-duration US Treasury bonds"],["LQD","US dollar investment-grade corporate bonds"]],[1,5.9])
p("Adjusted closes provide an approximate total-return series incorporating the vendor's dividend and split adjustments. They are not raw executable prices. The model treats changes in adjusted prices as changes in reinvested portfolio units. Fund operating expenses already affect the observed series and are not deducted a second time. Taxes and investor-specific dividend treatment are excluded.")
p("The downloader rejects empty responses, missing observations, duplicate dates, nonpositive prices and inconsistent symbol sets. It does not replace failed downloads with random values or forward-fill prices. Individual provider files and the consolidated series are saved with checksums and a download manifest. The research can then be rerun offline from this snapshot.")
h("Related work",2)
p("Moskowitz, Ooi and Pedersen [2] study time-series momentum across futures and forward contracts. Their work motivates examining past returns but does not validate this ETF rule: the present strategy ranks assets relative to one another and remains long only. Frazzini and Pedersen [3] study low-beta investing, which is related to defensive investing but differs from ranking ETFs on total historical volatility. Neither study is replicated here.")
p("Bailey and coauthors [4] explain the danger of selecting strategies after examining many backtests. Accordingly, this study reports its parameter grid as a sensitivity exercise and does not promote its best historical setting as a validated trading rule.")

page();h("Factor construction and portfolio rules")
h("Momentum and volatility",2)
p("For each ETF, daily return equals today's adjusted close divided by the previous session's adjusted close, minus one. Momentum is the analogous price ratio over 252 sessions. Volatility is the sample standard deviation of the last 252 daily returns, multiplied by the square root of 252. All observations used in a signal are available by that signal's closing date.")
p("Momentum = adjusted close today / adjusted close 252 sessions earlier - 1")
p("Annualized volatility = sample standard deviation of 252 daily returns × square root of 252")
p("The lookback uses trading observations rather than calendar days. There is no one-month momentum skip. A signal is eligible only when both inputs are available for every ETF in the configured universe.")
h("Combining comparable scores",2)
p("The original prototype directly subtracted volatility from momentum. Both are numerical rates, but their cross-sectional spreads differ, so equal coefficients do not ensure equal influence. The revised default first converts momentum and negative volatility into percentile ranks across the eight ETFs. A larger value is better for both transformed inputs.")
p("Composite score = 0.5 × momentum percentile rank + 0.5 × low-volatility percentile rank")
p("Average ranks are used for tied input values. If final scores tie, the configured ticker order breaks the tie deterministically. This transparent convention can affect a small universe and is not evidence that the tied assets have different expected returns. Ranks also discard information about the magnitude of differences, so normalization is a design choice rather than an assured performance improvement.")
h("Selection and execution",2)
p("At the final observed trading session of each month, the two highest-scoring ETFs receive target weights of 50% each. The signal is executed at the next observed trading close. Existing holdings earn that next session's price change before the trade; new holdings earn returns afterward. This one-session delay avoids assuming that a closing-price signal could have been known before trading at the same close.")
p("Holdings drift with prices until the following rebalance. The 50% allocation is a target at rebalance, not a continuous cap. There is no leverage, short position, stop loss, cash-switching rule, trend-sign filter or guaranteed downside floor. The strategy remains invested even when every ETF has negative momentum.")
p("The last monthly signal is stored for reproducibility even if its execution date lies beyond the sample. Such an unexecuted signal contributes no return to the reported backtest.")

page();h("Accounting and evaluation design")
h("Costs and self financing",2)
p("Every purchase and sale is charged 10 basis points of its traded dollar value. Selling a position and buying a replacement incurs costs on both legs. At each rebalance, the engine solves the fee against post-fee wealth before allocating the new target holdings. This prevents fees from creating an implicit borrowing balance. Initial purchases are charged; terminal liquidation is not assumed. Cash earns zero.")
p("For example, investing one unit of cash at a 0.1% purchase cost leaves 1 / 1.001 units invested. A flat market therefore produces a small initial loss, which is included in both performance and drawdown. The constant fee is a stylized all-in trading friction, not an estimate of every ETF's historical spread or market impact.")
h("Comparisons",2)
table(["Portfolio","Purpose"],[
 ["Rank blend","Default equal-weight combination of the two ranks"],
 ["Raw blend","Original momentum minus volatility score, corrected execution"],
 ["Momentum only","Isolates recent-return selection"],
 ["Low volatility only","Isolates defensive selection"],
 ["Equal weight universe","Monthly equal weights across the same eight ETFs"],
 ["SPY buy and hold","Single purchase of the equity benchmark"],
],[1.8,5.1])
p("All portfolios start from one unit of cash and trade at the same first evaluation close. The factor warm-up is excluded from every performance series. SPY measures equity opportunity cost; the equal-weight universe is a closer asset-coverage comparison. Neither benchmark is matched to the factor portfolio's risk exposure.")
h("Performance measures",2)
p("CAGR compounds final wealth over the number of observations divided by 252. Annualized volatility uses daily sample standard deviation. Sharpe is the annualized daily mean divided by annualized volatility, with an explicitly assumed zero risk-free rate. Maximum drawdown measures the largest decline from the running wealth peak, including initial capital. Annual traded notional sums purchases and sales and divides by sample years; it is not half-turnover.")
h("Verification",2)
p("Eleven deterministic tests cover signal timing, future-data independence, trading-calendar edges, price validation, price-driven weight drift, self-financing costs and metric definitions. A separate audit reconstructs every daily portfolio value from the persisted target orders and price ratios, verifies next-session execution, checks a direct SPY endpoint calculation and confirms that higher costs reduce each grid variant's CAGR.")

page();h("Full period results")
p(f"The common evaluation window contains {int(r.observations):,} sessions from {run['first_evaluation_date']} through {run['last_evaluation_date']}. Table 1 reports net results. The rank blend ends at {1+r.total_return:.3f} units for each initial unit, compared with {1+eq.total_return:.3f} for the equal-weight universe and {1+spy.total_return:.3f} for SPY.")
p("Table 1  Net performance over the full evaluation period", "Caption")
table(["Portfolio","CAGR","Volatility","Sharpe","Max drawdown"],[[name,pct(row.cagr),pct(row.annual_volatility),f"{row.sharpe_rf0:.2f}",pct(row.max_drawdown)] for name,row in s.iterrows()],[2.15,1.05,1.15,1,1.55])
picture("equity_curve.png","Figure 1  Growth of one unit after trading costs. All portfolios share the same start and end dates.")
p(f"Relative to momentum alone, the rank blend has a higher CAGR ({pct(r.cagr)} versus {pct(s.loc['Momentum only','cagr'])}) and lower volatility ({pct(r.annual_volatility)} versus {pct(s.loc['Momentum only','annual_volatility'])}). This supports a narrow descriptive benefit from combining the inputs within this universe and period.")
p(f"The broader claim fails: the rank blend trails the equal-weight universe in CAGR and Sharpe, with higher volatility and a larger maximum drawdown. Against SPY it reduces measured volatility and drawdown but gives up substantial growth. The raw blend earns {pct(s.loc['Raw blend','cagr'])}, above the rank blend's CAGR, while its Sharpe is lower. Rank normalization therefore does not dominate the original score.")

page();h("Drawdowns and changing conditions")
picture("drawdown.png","Figure 2  Net drawdowns relative to each portfolio's own previous peak, including initial capital.")
p(f"The default portfolio's maximum drawdown is {pct(r.max_drawdown)}. Low-volatility selection therefore does not establish a loss floor. In 2022 the rank blend loses {pct(-annual.loc[2022,'Rank blend'])}, compared with a {pct(-annual.loc[2022,'SPY buy and hold'])} loss for SPY. The simple intuition of protecting the downside is not reliable in every calendar year.")
p("Table 2  Descriptive subperiod results for the default and benchmarks", "Caption")
rows=[]
for name in ["Rank blend","Equal weight universe","SPY buy and hold"]:
    for period in ["2012-2019","2020-2025"]:
        row=periods[(periods.strategy==name)&(periods.period==period)].iloc[0]
        rows.append([name,period,pct(row.cagr),f"{row.sharpe_rf0:.2f}"])
table(["Portfolio","Period","CAGR","Sharpe"],rows,[2.7,1.5,1.3,1.4])
p("Subperiod statistics use the actual continuous strategy returns within each interval; portfolios are not reset or selected again at the split. Drawdowns for a subperiod, available in the exported table, are measured from wealth rebased at that interval's start. The labels are descriptive and must not be interpreted as an untouched validation set.")
lv1=periods[(periods.strategy=="Low volatility only")&(periods.period=="2012-2019")].iloc[0]
lv2=periods[(periods.strategy=="Low volatility only")&(periods.period=="2020-2025")].iloc[0]
p(f"The low-volatility-only portfolio changes markedly across the two intervals: its CAGR is {pct(lv1.cagr)} in 2012-2019 and {pct(lv2.cagr)} in 2020-2025. This variation cautions against treating a defensive historical characteristic as a stable return advantage. No causal attribution to a particular economic event is estimated here.")

page();h("Sensitivity to modeling choices")
p("The fixed grid varies the lookback across 126 and 252 sessions, the portfolio size across two and three ETFs, the momentum weight across 0.25, 0.50 and 0.75, and costs across 0, 10 and 25 basis points. All 36 combinations are exported. These cases are robustness checks; the default remains the equal-weight rank blend with a 252-session lookback and two positions.")
p("Table 3  CAGR across signal settings at 10 basis points per traded dollar", "Caption")
g10=grid[grid.cost_bps==10]
rows=[]
for lookback in [126,252]:
    for n in [2,3]:
        vals=[]
        for a in [.25,.5,.75]:
            val=g10[(g10.lookback==lookback)&(g10.top_n==n)&(g10.momentum_weight==a)].iloc[0]
            vals.append(pct(val.cagr))
        rows.append([lookback,n,*vals])
table(["Lookback","Positions","25% momentum","50% momentum","75% momentum"],rows,[1.1,1,1.6,1.6,1.6])
p("Table 4  Trading cost sensitivity of the default signal", "Caption")
gc=grid[(grid.lookback==252)&(grid.top_n==2)&(grid.momentum_weight==.5)].sort_values("cost_bps")
table(["Cost in bps","CAGR","Sharpe","Max drawdown"],[[int(x.cost_bps),pct(x.cagr),f"{x.sharpe_rf0:.2f}",pct(x.max_drawdown)] for x in gc.itertuples()],[1.6,1.6,1.5,2.2])
p(f"Across the entire grid, observed CAGR ranges from {pct(grid.cagr.min())} to {pct(grid.cagr.max())}. At the default signal settings, raising costs from zero to 25 basis points reduces CAGR from {pct(gc.iloc[0].cagr)} to {pct(gc.iloc[-1].cagr)}. The default portfolio trades approximately {r.annual_traded_notional:.2f} times portfolio wealth per year in combined purchases and sales, compared with {eq.annual_traded_notional:.2f} for the equal-weight universe. Its extra selection activity creates a meaningful cost burden.")
p("The grid is not a search for a winning headline. Even an attractive setting would be a retrospective finding requiring a genuinely new sample. The experiment does not adjust for multiple testing, estimate a probability of backtest overfitting, or calculate confidence intervals for differences between portfolios. The evidence is descriptive rather than a statistical claim of alpha.")
h("What the comparison establishes",2)
p("Combining factors can change the risk-return profile without beating simpler alternatives. Normalization changes which assets are selected, costs affect compounded outcomes, and the evaluation window changes how a strategy looks. These implementation choices are part of the research question rather than details to hide after selecting a favorable result.")

page();h("Limitations and conclusions")
h("Limits of the evidence",2)
p("The universe consists of eight surviving ETFs carried forward from the prototype. It is neither a historical census nor a point-in-time screen of all available funds. Selection and survivorship bias remain. The model has no delisting treatment or dynamic eligibility rules; adding newer funds would require explicit inception and missing-data policies.")
p("The adjusted-close snapshot comes from one public data provider and has not been reconciled against an independent source. Provider revisions can change later downloads. Complete observations and checksums improve reproducibility, but they do not certify economic correctness. An actual execution system would require raw prices, distribution accounting and instrument-specific trade constraints.")
p("Only a constant trading friction is modeled. Historical bid-ask spreads, taxes, market impact, liquidity constraints and order size are omitted. Cash earns no interest, and Sharpe uses a zero risk-free rate. Ranking on volatility does not manage correlations or guarantee diversification; two selected ETFs may share substantial exposures.")
p("The test begins after the 2008 financial crisis, so that episode is not represented. Although the sample includes multiple later conditions, it is limited to one realized market path. All observations were historical when this revision was developed. A split at 2020 is not a prospective test, and neither the full sample nor the sensitivity grid establishes live-trading readiness.")
h("Conclusion",2)
p(f"The original idea survives as a testable hypothesis, not as a confirmed claim of market outperformance. Over this historical sample, the equal-weight rank blend earns {pct(r.cagr)} annually after the stated costs and experiences a {pct(-r.max_drawdown)} peak-to-trough loss. It improves selected metrics relative to momentum-only selection, but it does not outperform the equal-weight universe or SPY on CAGR or the reported Sharpe measure. The low-volatility component changes exposure; it does not insure the portfolio.")
p("The project is stronger because its inputs, rules and limitations can now be inspected and its results reproduced. A sensible next research stage would freeze the present specification before observing new data, evaluate a future paper-trading period, and reconcile market data independently. Those steps are proposed future work, not completed achievements.")
h("Revision provenance",2)
p("The original four-script prototype and original paper are retained in the project archive. This October 2026 revision was prepared with AI assistance for implementation, testing, analysis and writing. The historical experiment, accounting corrections, comparison portfolios and updated findings belong to this revision. The paper makes no claim that these additions were completed during the original capstone or that the student independently authored every new component.")

page();h("Reproducibility and references")
h("Research package",2)
p("The accompanying folder includes the original materials, configuration, documented source code, a verified price snapshot, daily returns, actual weights, trade ledgers, annual returns, subperiod tables, the full sensitivity grid, tests and this editable paper. The run manifest records package versions and hashes of both the data and research code.")
table(["Item","Recorded value"],[
 ["Data source","Yahoo Finance through yfinance"],
 ["Snapshot dates",f"{manifest['first_observation']} to {manifest['last_observation']}"],
 ["Evaluation dates",f"{run['first_evaluation_date']} to {run['last_evaluation_date']}"],
 ["Universe and evaluation rows",f"8 ETFs; {int(r.observations):,} daily rows per portfolio"],
 ["yfinance and pandas",f"{run['packages']['yfinance']} and {run['packages']['pandas']}"],
 ["Default cost","10 bps per dollar purchased or sold"],
 ["Audit result",audit["status"]],
],[2.35,4.55])
p("To reproduce results, install the dependencies, run research.py run against the saved data, execute the tests, and run audit_results.py. The download command is separate so a later vendor revision does not silently replace the submitted snapshot. build_paper.py reads the exported metrics; the document can be regenerated after a new experiment. The README contains exact commands and definitions.")
h("References",2)
refs=[
 "[1] yfinance. Download API documentation. Accessed October 7, 2026. https://ranaroussi.github.io/yfinance/reference/api/yfinance.download.html",
 "[2] Moskowitz, T. J., Ooi, Y. H., and Pedersen, L. H. (2012). Time series momentum. Journal of Financial Economics, 104(2), 228-250. https://doi.org/10.1016/j.jfineco.2011.11.003",
 "[3] Frazzini, A., and Pedersen, L. H. (2014). Betting against beta. Journal of Financial Economics, 111(1), 1-25. https://doi.org/10.1016/j.jfineco.2013.10.005",
 "[4] Bailey, D. H., Borwein, J. M., Lopez de Prado, M., and Zhu, Q. J. The probability of backtest overfitting. Author-hosted manuscript. https://www.davidhbailey.com/dhbpapers/backtest-prob.pdf",
 "[5] Chen, C. ETF Multi-Factor Strategy. Original local four-script prototype and capstone paper, preserved in original/. Current revised source, data manifest and results accompany this paper.",
]
for ref in refs:
    para=p(ref)
    para.paragraph_format.space_after=Pt(8)
    for rr in para.runs: rr.font.size=Pt(10)
p("Sources [2]-[4] provide conceptual context and methodological cautions. All portfolio statistics in this paper are computed from the accompanying local experiment; none are performance claims borrowed from those publications.")

doc.save(OUT/"ETF_Research_Revised.docx")
(OUT/"ETF_Research_Revised.md").write_text("\n".join(md))
print(OUT/"ETF_Research_Revised.docx")
